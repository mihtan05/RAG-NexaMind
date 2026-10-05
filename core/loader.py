"""
Module: core/loader.py
Chức năng: Nạp tài liệu đa định dạng (PDF, Word, Excel, PowerPoint, TXT, HTML...)
sử dụng Microsoft MarkItDown làm bộ chuyển đổi chính, kèm các lớp Fallback tự động
(mammoth, python-docx, pdfplumber, openpyxl) để đảm bảo không bao giờ bị gián đoạn.
"""

import os
from pathlib import Path
from typing import Dict, Any, Optional
from markitdown import MarkItDown


class DocumentLoader:
    """
    Trình nạp và chuyển đổi tài liệu đa định dạng thành văn bản Markdown chuẩn:
    1. Ưu tiên sử dụng Microsoft MarkItDown.
    2. Tự động kích hoạt cơ chế Fallback (mammoth / python-docx / openpyxl / pypdf) nếu
       MarkItDown gặp lỗi thiếu dependency hoặc lỗi định dạng.
    """

    SUPPORTED_EXTENSIONS = {
        ".pdf", ".docx", ".doc", ".xlsx", ".xls",
        ".pptx", ".ppt", ".html", ".htm", ".txt",
        ".csv", ".json", ".xml", ".md"
    }

    def __init__(self):
        """Khởi tạo phiên bản MarkItDown."""
        self.md = MarkItDown()

    def is_supported(self, file_path: str) -> bool:
        """Kiểm tra định dạng file có được hỗ trợ hay không."""
        ext = Path(file_path).suffix.lower()
        return ext in self.SUPPORTED_EXTENSIONS

    def _convert_docx_fallback(self, file_path: str) -> str:
        """Fallback chuyển đổi Word (.docx) sang Markdown bằng mammoth hoặc python-docx."""
        # 1. Thử mammoth (chuyển sang markdown rất đẹp)
        try:
            import mammoth
            with open(file_path, "rb") as docx_file:
                result = mammoth.convert_to_markdown(docx_file)
                if result and result.value and result.value.strip():
                    return result.value.strip()
        except Exception:
            pass

        # 2. Thử python-docx
        try:
            import docx
            doc = docx.Document(file_path)
            parts = []
            for p in doc.paragraphs:
                text = p.text.strip()
                if not text:
                    continue
                style_name = p.style.name.lower() if p.style else ""
                if "heading 1" in style_name:
                    parts.append(f"# {text}")
                elif "heading 2" in style_name:
                    parts.append(f"## {text}")
                elif "heading 3" in style_name:
                    parts.append(f"### {text}")
                else:
                    parts.append(text)

            # Đọc thêm bảng biểu trong file docx
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
                    parts.append("| " + " | ".join(row_cells) + " |")

            return "\n\n".join(parts).strip()
        except Exception:
            pass

        return ""

    def _convert_pdf_fallback(self, file_path: str) -> str:
        """Fallback chuyển đổi PDF sang văn bản Markdown."""
        # Thử pdfplumber
        try:
            import pdfplumber
            parts = []
            with pdfplumber.open(file_path) as pdf:
                for page_idx, page in enumerate(pdf.pages, 1):
                    text = page.extract_text()
                    if text and text.strip():
                        parts.append(f"### Trang {page_idx}\n{text.strip()}")
            if parts:
                return "\n\n".join(parts)
        except Exception:
            pass

        # Thử pypdf
        try:
            import pypdf
            reader = pypdf.PdfReader(file_path)
            parts = []
            for page_idx, page in enumerate(reader.pages, 1):
                text = page.extract_text()
                if text and text.strip():
                    parts.append(f"### Trang {page_idx}\n{text.strip()}")
            if parts:
                return "\n\n".join(parts)
        except Exception:
            pass

        return ""

    def _convert_tabular_fallback(self, file_path: str) -> str:
        """Fallback chuyển đổi Excel (.xlsx, .xls) sang Markdown bảng."""
        try:
            import pandas as pd
            excel_data = pd.read_excel(file_path, sheet_name=None)
            parts = []
            for sheet_name, df in excel_data.items():
                parts.append(f"### Sheet: {sheet_name}\n")
                parts.append(df.to_markdown(index=False))
            return "\n\n".join(parts)
        except Exception:
            pass
        return ""

    def load_document(self, file_path: str) -> Dict[str, Any]:
        """
        Nạp một file từ đường dẫn trên đĩa và chuyển đổi sang Markdown.

        Args:
            file_path: Đường dẫn tuyệt đối hoặc tương đối tới tài liệu.

        Returns:
            dict chứa: filename, content, extension, filepath.
        """
        path_obj = Path(file_path)
        if not path_obj.exists():
            raise FileNotFoundError(f"Không tìm thấy file: {file_path}")

        ext = path_obj.suffix.lower()
        text_content = ""

        # Bước 1: Thử dùng Microsoft MarkItDown với instance tươi
        try:
            md_engine = MarkItDown()
            result = md_engine.convert(str(path_obj.resolve()))
            if result and result.text_content and result.text_content.strip():
                text_content = result.text_content.strip()
        except Exception:
            text_content = ""

        # Bước 2: Tự động Fallback nếu MarkItDown không đọc được
        if not text_content:
            if ext in {".docx", ".doc"}:
                text_content = self._convert_docx_fallback(str(path_obj.resolve()))
            elif ext == ".pdf":
                text_content = self._convert_pdf_fallback(str(path_obj.resolve()))
            elif ext in {".xlsx", ".xls"}:
                text_content = self._convert_tabular_fallback(str(path_obj.resolve()))
            elif ext in {".txt", ".md", ".csv", ".json", ".xml"}:
                try:
                    with open(path_obj, "r", encoding="utf-8", errors="ignore") as f:
                        text_content = f.read().strip()
                except Exception:
                    text_content = ""

        if not text_content:
            raise RuntimeError(
                f"Không thể trích xuất nội dung từ file {path_obj.name}. "
                f"Vui lòng kiểm tra lại định dạng hoặc nội dung của file."
            )

        return {
            "filename": path_obj.name,
            "content": text_content,
            "extension": ext,
            "filepath": str(path_obj.resolve()),
        }

    def load_uploaded_file(self, uploaded_file, save_dir: str = "data") -> Dict[str, Any]:
        """
        Hỗ trợ nạp file tải lên từ Streamlit (UploadedFile), lưu tạm vào thư mục
        và gọi load_document.
        """
        os.makedirs(save_dir, exist_ok=True)
        file_path = os.path.join(save_dir, uploaded_file.name)

        # Lưu nội dung buffer ra đĩa
        with open(file_path, "wb") as f:
            f.write(uploaded_file.getbuffer())

        return self.load_document(file_path)
